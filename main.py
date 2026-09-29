from pathlib import Path
import csv
from statistics import mean, median, stdev

from openpyxl import Workbook, load_workbook


def create_excel_from_csvs(
    data_path: Path,
    csv_pattern: str,
    output_filename: str
):
    output_path = data_path / output_filename

    workbook = Workbook(write_only=True)

    csv_files = sorted(data_path.glob(csv_pattern))

    for csv_file in csv_files:
        sheet_name = csv_file.stem[:31]
        worksheet = workbook.create_sheet(title=sheet_name)

        with csv_file.open(
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:
            reader = csv.reader(file, delimiter=",")

            for row_number, row in enumerate(reader):
                if row_number == 0:
                    worksheet.append(row)
                else:
                    converted_row = []
                    for value in row:
                        try:
                            converted_value = int(value)
                        except ValueError:
                            converted_value = float(value)
                        converted_row.append(converted_value)
                    worksheet.append(converted_row)
                row_number += 1

    workbook.save(output_path)

    print(f"Erstellt: {output_path}")


def calculate_statistics(values):
    return {
        "MIN": min(values),
        "AVERAGE": mean(values),
        "MEDIAN": median(values),
        "MAX": max(values),
        "STDEV": stdev(values),
    }

def add_statistics_to_sheet(sheet):
    last_row = sheet.max_row
    last_column = sheet.max_column

    statistics = {}

    for column in range(2, last_column + 1):
        values = []
        for row in range(2, last_row + 1):
            cell_value = sheet.cell(row=row, column=column).value
            if isinstance(cell_value, (int, float)):
                values.append(cell_value)

        if values:
            statistics[column] = calculate_statistics(values)

    statistics_start_row = last_row + 1

    for offset, statistic_name in enumerate(["MIN", "AVERAGE", "MEDIAN", "MAX", "STDEV"]):
        row = statistics_start_row + offset
        sheet.cell(
            row=row,
            column=1,
            value=statistic_name
        )
        for column, values in statistics.items():
            sheet.cell(
                row=row,
                column=column,
                value=values[statistic_name]
            )


def main():
    data_path = Path(
        r"C:\Users\ma1021716\IdeaProjects\RoulettePlus\data"
    )

    # Alle spieler_*.csv → spieler.xlsx
    create_excel_from_csvs(
        data_path,
        "spieler*.csv",
        "spieler.xlsx"
    )

    # Alle quoten_*.csv → quoten.xlsx
    create_excel_from_csvs(
        data_path,
        "quoten*.csv",
        "quoten.xlsx"
    )

    workbook = load_workbook(data_path / "quoten.xlsx")
    for sheet in workbook.worksheets:
        print(f"Berechne Statistiken für: {sheet.title}")
        add_statistics_to_sheet(sheet)
    workbook.save(data_path / "quoten.xlsx")

    print(f"Statistiken hinzugefügt: {data_path / 'quoten.xlsx'}")

if __name__ == "__main__":
    main()